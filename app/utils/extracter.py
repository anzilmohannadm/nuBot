import json
import openpyxl
from llama_index.core.readers.base import BaseReader
from llama_index.core import Document


class xlsReader(BaseReader):
    def load_data(self, file, extra_info=None):
        with open(file, "r") as f:
            text = f.read()
        # load_data returns a list of Document objects
        return [Document(text=text, extra_info=extra_info or {})]
        

class xsdReader(BaseReader):
    def load_data(self, file, extra_info=None):
        with open(file, "r") as f:
            text = f.read()
        # load_data returns a list of Document objects
        return [Document(text=text, extra_info=extra_info or {})]

class xlsxReader(BaseReader):
    def load_data(self, file, extra_info=None):
        with openpyxl.load_workbook(file) as workbook:
            # Select the active worksheet
            sheet = workbook.active
            # Initialize an empty string to store the data
            text = ""
            # Extract data from the worksheet and format as a string table
            for row in sheet.iter_rows(values_only=True):
                # Join the row elements into a string with tab separation
                text += '\t'.join(str(cell) if cell is not None else '' for cell in row) + '\n'
        
        return [Document(text=text, extra_info=extra_info or {})]



file_extracter = {"xls":xlsReader(),"xlsx":xlsxReader()}

class JSONReader(BaseReader):
    def load_data(self, file, extra_info=None):
        with open(file, "r") as f:
            data = json.load(f)
        # Assuming the JSON has a 'text' field and other metadata
        text = data.get("text", "")
        metadata = {k: v for k, v in data.items() if k != "text"}
        return [Document(text=text, extra_info=metadata)]


file_extracter = {"xls":xlsReader(),"xlsx":xlsxReader(),"json":JSONReader()}